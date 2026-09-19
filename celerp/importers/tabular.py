# Copyright (c) 2026 Noah Severs
# SPDX-License-Identifier: LicenseRef-Proprietary

"""Tabular file parsing shared by the CSV import UI and the agent importer.

CSV helpers (column mapping, cell validation, error reports) are defined here
once and re-imported by ``ui/routes/csv_import.py`` so both the UI and the
capability agent read the same source of truth. ``read_csv``/``read_xlsx``/
``read_table`` turn an uploaded file into ``(header, rows)`` with the same
string shape whatever the source format, so a single ``validate_cell`` applies.
"""

from __future__ import annotations

import csv
import datetime
import io
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from ui.i18n import t


ValidateFn = Callable[[str, str, dict], bool]

# Size guards. MAX_XLSX_BYTES matches the upload bound; the uncompressed and
# entry-count guards stop a zip bomb before openpyxl ever parses the workbook.
MAX_ROWS = 10_000
MAX_CELLS = 200_000
MAX_XLSX_BYTES = 10 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_XLSX_ENTRIES = 4096


class TabularError(ValueError):
    """A file that cannot be turned into a table. Carries the offending cell
    position (``row``/``column``) for formula errors and the available sheet
    names (``sheets``) when the caller must choose one."""

    def __init__(
        self,
        message: str,
        *,
        row: int | None = None,
        column: str | None = None,
        sheets: list[str] | None = None,
    ) -> None:
        super().__init__(message)
        self.row = row
        self.column = column
        self.sheets = sheets


# Columns always shown in the error table (identifiers), even if they have no errors.
_IDENTIFIER_COLS = {"sku", "name", "id", "email", "code"}


@dataclass(frozen=True)
class CsvImportSpec:
    cols: list[str]
    required: set[str]
    type_map: dict[str, Callable[[str], Any]]


def validate_cell(spec: CsvImportSpec, col: str, value: str, row: dict | None = None) -> bool:
    if col in spec.required and not value.strip():
        return False
    cast = spec.type_map.get(col)
    if cast and value.strip():
        try:
            cast(value)
        except (ValueError, TypeError):
            return False
    return True


# ---------------------------------------------------------------------------
# Column mapping
# ---------------------------------------------------------------------------

# Common aliases: CSV header (lowercase) -> Celerp target field.
# Used to pre-fill the mapping dropdown. Not auto-committed - user always sees
# and confirms the suggestion.
_COMMON_ALIASES: dict[str, str] = {
    "item_type": "category",
    "type": "category",
    "product_type": "category",
    "price": "retail_price",
    "selling_price": "retail_price",
    "sale_price": "retail_price",
    "cost": "cost_price",
    "unit_cost": "cost_price",
    "purchase_price": "cost_price",
    "total_cost": "cost_price_total",
    "cost_total": "cost_price_total",
    "total cost": "cost_price_total",
    "cost total": "cost_price_total",
    "wholesale": "wholesale_price",
    "weight_ct": "weight",
    "weight_g": "weight",
    "location": "location_name",
    "warehouse": "location_name",
    "upc": "barcode",
    "ean": "barcode",
    "isbn": "barcode",
    "code": "sku",
    "item_code": "sku",
    "product_code": "sku",
    "product_name": "name",
    "item_name": "name",
    "title": "name",
    "desc": "description",
    "qty": "quantity",
    "stock": "quantity",
    "on_hand": "quantity",
    # spaced variants that don't exact-match underscore targets
    "sell by": "sell_by",
    "purchase unit": "purchase_unit",
    "weight unit": "weight_unit",
    "location name": "location_name",
    "hs code": "hs_code",
    "purchase sku": "purchase_sku",
    "purchase name": "purchase_name",
    "purchase conversion factor": "purchase_conversion_factor",
    "short description": "short_description",
}

# Aliases for category attribute keys (csv_col_lower → attr_key_lower).
# Used in suggest_mapping Pass 2b to bridge common spreadsheet column names
# to their canonical category attribute counterparts.
_COMMON_ATTR_ALIASES: dict[str, str] = {
    "stone_color": "color",
    "stone_colour": "color",
    "stone_shape": "shape",
    "stone_treatment": "treatment",
    "stone_origin": "origin",
    "color_grade": "grade",
    "colour_grade": "grade",
    "clarity_grade": "clarity",
    "certificate_number": "certificate_no",
    "cert_number": "certificate_no",
    "cert_no": "certificate_no",
}

# Columns that should always default to Skip (system-managed; never imported)
_FORCE_SKIP_COLS: frozenset[str] = frozenset({"created_at", "updated_at", "status"})

# Sentinel values for the mapping dropdown
MAPPING_ATTRIBUTE = "__attr__"
MAPPING_SKIP = "__skip__"
MAPPING_ATTR_PREFIX = "__catattr:"  # Category attribute: "__catattr:stone_type"


def suggest_mapping(
    csv_cols: list[str],
    target_cols: list[str],
    category_attrs: list[str] | None = None,
) -> dict[str, str]:
    """Return {csv_col: suggested_target} for each CSV column.

    Priority:
    0. Force-skip columns (created_at, updated_at, status) → always MAPPING_SKIP
    1. Exact match (case-insensitive) to a core target column
    2. Known alias match to a core target column
    2b. Known alias match to a category attribute key
    3. Exact match to a category attribute key (prefixed with MAPPING_ATTR_PREFIX)
    4. Default to MAPPING_ATTRIBUTE (import as custom field)

    Each target field is claimed at most once (first match wins).
    """
    mapping: dict[str, str] = {}
    claimed: set[str] = set()
    target_lower = {item.lower(): item for item in target_cols}
    attrs = category_attrs or []
    attr_lower = {a.lower().replace(" ", "_"): a for a in attrs}

    # Pass 0: force-skip system columns
    for csv_col in csv_cols:
        if csv_col.lower().strip() in _FORCE_SKIP_COLS:
            mapping[csv_col] = MAPPING_SKIP

    # Pass 1: exact matches to core fields (also try space→underscore normalization)
    for csv_col in csv_cols:
        if csv_col in mapping:
            continue
        lc = csv_col.lower().strip()
        lc_norm = lc.replace(" ", "_")
        match = target_lower.get(lc) or target_lower.get(lc_norm)
        if match and match not in claimed:
            mapping[csv_col] = match
            claimed.add(match)

    # Pass 2: alias matches to core fields
    for csv_col in csv_cols:
        if csv_col in mapping:
            continue
        lc = csv_col.lower().strip()
        alias_target = _COMMON_ALIASES.get(lc)
        if alias_target and alias_target in target_lower.values() and alias_target not in claimed:
            mapping[csv_col] = alias_target
            claimed.add(alias_target)

    # Pass 2b: alias matches to category attribute keys
    claimed_attrs: set[str] = set()
    for csv_col in csv_cols:
        if csv_col in mapping:
            continue
        lc = csv_col.lower().strip().replace(" ", "_")
        alias_attr = _COMMON_ATTR_ALIASES.get(lc)
        if alias_attr and alias_attr in attr_lower.values() and alias_attr not in claimed_attrs:
            mapping[csv_col] = f"{MAPPING_ATTR_PREFIX}{alias_attr}"
            claimed_attrs.add(alias_attr)

    # Pass 3: exact match to category attribute keys
    for csv_col in csv_cols:
        if csv_col in mapping:
            continue
        lc = csv_col.lower().strip().replace(" ", "_")
        if lc in attr_lower and attr_lower[lc] not in claimed_attrs:
            mapping[csv_col] = f"{MAPPING_ATTR_PREFIX}{attr_lower[lc]}"
            claimed_attrs.add(attr_lower[lc])

    # Pass 4: everything else defaults to custom
    for csv_col in csv_cols:
        if csv_col not in mapping:
            mapping[csv_col] = MAPPING_ATTRIBUTE

    return mapping


def validate_column_mapping(
    form: dict,
    csv_cols: list[str],
    core_fields: set[str] | None = None,
) -> list[str]:
    """Validate the user's column mapping choices. Returns list of error messages (empty = valid).

    Checks:
    1. Two CSV columns mapped to the same target field (duplicate targets).
    2. Attribute names that collide with core/built-in field names.
    3. Two attribute columns with the same custom name.
    """
    errors: list[str] = []
    core = core_fields or set()

    # Collect all mappings
    target_sources: dict[str, list[str]] = {}  # target -> [csv_col, ...]
    attr_names: dict[str, list[str]] = {}  # attr_name -> [csv_col, ...]

    for col in csv_cols:
        target = str(form.get(f"map__{col}", MAPPING_ATTRIBUTE) or MAPPING_ATTRIBUTE)
        if target == MAPPING_SKIP:
            continue

        if target == MAPPING_ATTRIBUTE:
            # Custom field name (from text input) or original col name
            attr_name = str(form.get(f"attr_name__{col}", "") or "").strip() or col
            attr_names.setdefault(attr_name, []).append(col)
            # Check collision with core field names
            if attr_name.lower() in {c.lower() for c in core}:
                errors.append(
                    t("import.err_custom_name_conflict", name=attr_name, col=col)
                )
        elif target.startswith(MAPPING_ATTR_PREFIX):
            # Category attribute - use the attr key as the attribute name
            attr_key = target[len(MAPPING_ATTR_PREFIX):]
            attr_names.setdefault(attr_key, []).append(col)
        else:
            target_sources.setdefault(target, []).append(col)

    # Check duplicate target fields
    for target, sources in target_sources.items():
        if len(sources) > 1:
            names = " and ".join(f'"{s}"' for s in sources)
            errors.append(
                t(
                    "import.err_duplicate_target",
                    cols=names,
                    target=target.replace("_", " ").title(),
                )
            )

    # Check duplicate attribute names
    for attr_name, sources in attr_names.items():
        if len(sources) > 1:
            names = " and ".join(f'"{s}"' for s in sources)
            errors.append(
                t("import.err_duplicate_attr", cols=names, name=attr_name)
            )

    return errors


def mapped_field_name(col: str, target: str, attr_name: str | None = None) -> str | None:
    """Destination row key for a mapped column, or ``None`` to drop it.

    The single source of the mapping semantics shared by the browser importer
    (``apply_column_mapping``) and the agent importer: ``MAPPING_SKIP`` drops the
    column; ``MAPPING_ATTRIBUTE`` keeps it as a custom attribute under
    ``attr_name`` (falling back to the original header); ``MAPPING_ATTR_PREFIX``
    uses the category-attribute key after the prefix; anything else is the core
    target field name.
    """
    if target == MAPPING_SKIP:
        return None
    if target == MAPPING_ATTRIBUTE:
        return attr_name or col
    if target.startswith(MAPPING_ATTR_PREFIX):
        return target[len(MAPPING_ATTR_PREFIX):]
    return target


def remap_rows(
    cols: list[str],
    rows: list[dict],
    mapping: dict[str, str],
    attr_names: dict[str, str] | None = None,
) -> tuple[list[str], list[dict]]:
    """Apply a ``{col: target}`` mapping to already-parsed rows.

    Preserves the original column order, drops ``MAPPING_SKIP`` columns, and
    renames the rest via :func:`mapped_field_name`. Returns ``(new_cols, rows)``.
    """
    attr_names = attr_names or {}
    new_cols: list[str] = []
    rename: dict[str, str] = {}
    for col in cols:
        dest = mapped_field_name(col, mapping.get(col, MAPPING_ATTRIBUTE), attr_names.get(col))
        if dest is None:
            continue
        new_cols.append(dest)
        rename[col] = dest
    remapped = [{rename[c]: row.get(c, "") for c in rename} for row in rows]
    return new_cols, remapped


def apply_column_mapping(form: dict, csv_text: str) -> tuple[str, list[str]]:
    """Apply user's column mapping to CSV data.

    Reads map__<col>=<target> fields from the form.  Renames CSV headers
    according to the mapping.  Columns mapped to MAPPING_SKIP are dropped.
    Columns mapped to MAPPING_ATTRIBUTE use the custom name from attr_name__<col>
    (falling back to the original header).

    Returns (remapped_csv_text, remapped_cols).
    """
    reader = csv.DictReader(io.StringIO(csv_text))
    original_cols = list(reader.fieldnames or [])
    rows = list(reader)

    # Parse mapping from form
    mapping: dict[str, str] = {}
    for col in original_cols:
        target = str(form.get(f"map__{col}", MAPPING_ATTRIBUTE) or MAPPING_ATTRIBUTE)
        mapping[col] = target

    # Build new column list and rename map
    new_cols: list[str] = []
    rename: dict[str, str] = {}  # original -> new name
    for col in original_cols:
        attr_name = str(form.get(f"attr_name__{col}", "") or "").strip() or None
        dest = mapped_field_name(col, mapping[col], attr_name)
        if dest is None:
            continue
        new_cols.append(dest)
        rename[col] = dest

    # Write remapped CSV
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=new_cols, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        new_row = {}
        for col in original_cols:
            if col not in rename:
                continue
            new_row[rename[col]] = row.get(col, "")
        writer.writerow(new_row)

    return output.getvalue(), new_cols


def apply_fixes_to_rows(
    form: dict,
    rows: list[dict],
    cols: list[str],
) -> list[dict]:
    """Apply inline-fix form values back into the row dicts.

    Reads ``fixes_json``: a JSON object ``{"row__col": value, ...}`` serialized
    by the fix form's submit handler. One field regardless of error count,
    so Starlette's max_fields limit is never hit.
    """
    import json as _json

    fixes_raw = form.get("fixes_json", "")
    if not fixes_raw:
        return rows
    try:
        fixes: dict = _json.loads(fixes_raw)
    except (ValueError, TypeError):
        return rows

    cols_set = set(cols)
    for key, value in fixes.items():
        if not isinstance(key, str):
            continue
        parts = key.split("__", 1)
        if len(parts) != 2:
            continue
        ri_str, col = parts
        try:
            ri = int(ri_str)
        except ValueError:
            continue
        if 0 <= ri < len(rows) and col in cols_set:
            rows[ri][col] = str(value)
    return rows


def _row_errors(row: dict, cols: list[str], validate: ValidateFn) -> list[str]:
    return [col for col in cols if not validate(col, str(row.get(col, "")), row)]


def error_report_csv(rows: list[dict], cols: list[str], validate: ValidateFn) -> str:
    """Return CSV with original columns + an `_errors` column, containing only invalid rows."""
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=cols + ["_errors"], extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        bad = _row_errors(row, cols, validate)
        if bad:
            writer.writerow({**{c: row.get(c, "") for c in cols}, "_errors": "; ".join(bad)})
    return output.getvalue()


def _rows_to_csv(rows: list[dict], cols: list[str]) -> str:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=cols, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue()


# ---------------------------------------------------------------------------
# File readers
# ---------------------------------------------------------------------------


def _enforce_bounds(cols: list[str], rows: list[dict]) -> None:
    if len(rows) > MAX_ROWS:
        raise TabularError(f"Too many rows: {len(rows)} exceeds the {MAX_ROWS} limit.")
    cells = len(rows) * max(len(cols), 1)
    if cells > MAX_CELLS:
        raise TabularError(f"Too many cells: {cells} exceeds the {MAX_CELLS} limit.")


def read_csv(text: str) -> tuple[list[str], list[dict]]:
    """Parse CSV text into (header, rows) via csv.DictReader, BOM stripped and
    row/cell bounds enforced."""
    if text.startswith("﻿"):
        text = text[1:]
    reader = csv.DictReader(io.StringIO(text))
    cols = list(reader.fieldnames or [])
    rows = list(reader)
    _enforce_bounds(cols, rows)
    return cols, rows


def _stringify(value: Any) -> str:
    """Render an xlsx cell value the way the CSV path sees it: integral numbers
    without a trailing ``.0``, dates as ISO, everything else as its string."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime.datetime):
        if (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0):
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)


def _sheet_is_empty(ws) -> bool:
    for row in ws.iter_rows(values_only=True):
        if any(c is not None and str(c).strip() != "" for c in row):
            return False
    return True


def read_xlsx(data: bytes, *, sheet: str | None) -> tuple[list[str], list[dict]]:
    """Parse an .xlsx workbook into (header, rows).

    A zipfile pre-check bounds the entry count and total uncompressed size
    before openpyxl parses anything, so a zip bomb is rejected up front. Exactly
    one non-empty sheet auto-selects; several with no chosen sheet raise with the
    available names. A formula cell raises with its position.
    """
    if len(data) > MAX_XLSX_BYTES:
        raise TabularError("File is too large.")

    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise TabularError("File is not a valid .xlsx workbook.") from exc
    infos = archive.infolist()
    if len(infos) > MAX_XLSX_ENTRIES:
        raise TabularError("Workbook has too many internal entries.")
    if sum(info.file_size for info in infos) > MAX_XLSX_UNCOMPRESSED:
        raise TabularError("Workbook is too large when uncompressed.")

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except (InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
        raise TabularError("Could not open the workbook.") from exc

    try:
        names = list(workbook.sheetnames)
        if sheet is not None:
            if sheet not in names:
                raise TabularError(f"Sheet {sheet!r} is not in the workbook.", sheets=names)
            worksheet = workbook[sheet]
        else:
            non_empty = [name for name in names if not _sheet_is_empty(workbook[name])]
            if not non_empty:
                raise TabularError("The workbook has no data.")
            if len(non_empty) > 1:
                raise TabularError("Choose a sheet.", sheets=non_empty)
            worksheet = workbook[non_empty[0]]

        header: list[str] = []
        rows: list[dict] = []
        n_cols = 0
        for cells in worksheet.iter_rows():
            values: list[str] = []
            for cell in cells:
                value = cell.value
                if getattr(cell, "data_type", None) == "f" or (
                    isinstance(value, str) and value.startswith("=")
                ):
                    raise TabularError(
                        "Formula cells are not supported.",
                        row=cell.row,
                        column=getattr(cell, "column_letter", None),
                    )
                values.append(_stringify(value))
            if not header:
                header = values
                n_cols = len(header)
                continue
            if len(rows) >= MAX_ROWS:
                raise TabularError(f"Too many rows: exceeds the {MAX_ROWS} limit.")
            if (len(rows) + 1) * max(n_cols, 1) > MAX_CELLS:
                raise TabularError(f"Too many cells: exceeds the {MAX_CELLS} limit.")
            row = {header[i]: (values[i] if i < len(values) else "") for i in range(n_cols)}
            rows.append(row)
        return header, rows
    finally:
        workbook.close()


def read_table(
    data: bytes,
    filename: str,
    *,
    sheet: str | None = None,
) -> tuple[list[str], list[dict]]:
    """Dispatch on the file suffix. CSV and .xlsx are supported; .xlsm/.xls and
    everything else raise."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return read_csv(data.decode("utf-8-sig"))
    if suffix == ".xlsx":
        return read_xlsx(data, sheet=sheet)
    if suffix in (".xlsm", ".xls"):
        raise TabularError(f"{suffix} files are not supported; save as .xlsx or .csv.")
    raise TabularError(f"Unsupported file type: {suffix or filename!r}.")
