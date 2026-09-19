# Copyright (c) 2026 Noah Severs
# SPDX-License-Identifier: LicenseRef-Proprietary
"""Tabular reader: CSV parity with the old UI importer plus xlsx parsing."""
from __future__ import annotations

import io
import zipfile

import openpyxl
import pytest

from celerp.importers import tabular


def _xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    """Build an .xlsx workbook from {sheet_name: [[row cells], ...]}."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        worksheet = workbook.create_sheet(title=name)
        for row in rows:
            worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_csv_parity_bom_quoting_and_identifier_strings():
    """A leading BOM is stripped from the first header, quoted fields keep their
    embedded commas, and identifier-like values stay strings (never coerced)."""
    text = '﻿sku,name,note\r\n007,"Acme, Inc",hi\r\n'
    cols, rows = tabular.read_csv(text)
    assert cols == ["sku", "name", "note"]
    assert rows == [{"sku": "007", "name": "Acme, Inc", "note": "hi"}]


def test_xlsx_single_sheet_autoselect():
    data = _xlsx_bytes({"Only": [["sku", "qty"], ["A1", 3], ["B2", 4]]})
    cols, rows = tabular.read_xlsx(data, sheet=None)
    assert cols == ["sku", "qty"]
    assert rows == [{"sku": "A1", "qty": "3"}, {"sku": "B2", "qty": "4"}]


def test_xlsx_multiple_sheets_requires_choice():
    data = _xlsx_bytes({
        "First": [["sku"], ["A1"]],
        "Second": [["name"], ["Widget"]],
    })
    with pytest.raises(tabular.TabularError) as excinfo:
        tabular.read_xlsx(data, sheet=None)
    assert excinfo.value.sheets == ["First", "Second"]

    cols, rows = tabular.read_xlsx(data, sheet="Second")
    assert cols == ["name"]
    assert rows == [{"name": "Widget"}]


def test_xlsx_formula_cell_rejected_with_position():
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Calc"
    worksheet.append(["sku", "total"])
    worksheet["A2"] = "Widget"
    worksheet["B2"] = "=A2&A2"
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(tabular.TabularError) as excinfo:
        tabular.read_xlsx(buffer.getvalue(), sheet=None)
    assert excinfo.value.row == 2
    assert excinfo.value.column == "B"


def test_xlsm_xls_rejected():
    data = _xlsx_bytes({"Only": [["sku"], ["A1"]]})
    for filename in ("book.xlsm", "book.xls"):
        with pytest.raises(tabular.TabularError):
            tabular.read_table(data, filename)


def test_zip_bomb_rejected_before_parse():
    """A workbook whose entries expand past the uncompressed cap is rejected by
    the pre-check, never handed to openpyxl."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("payload.bin", b"\x00" * (tabular.MAX_XLSX_UNCOMPRESSED + 1))
    data = buffer.getvalue()
    assert len(data) < tabular.MAX_XLSX_BYTES

    with pytest.raises(tabular.TabularError) as excinfo:
        tabular.read_xlsx(data, sheet=None)
    assert "uncompressed" in str(excinfo.value).lower()


def test_row_and_cell_bounds():
    too_many_rows = "col\n" + "x\n" * (tabular.MAX_ROWS + 1)
    with pytest.raises(tabular.TabularError) as rows_err:
        tabular.read_csv(too_many_rows)
    assert "rows" in str(rows_err.value).lower()

    n_cols = 210
    n_rows = 1000  # under MAX_ROWS, but n_cols * n_rows exceeds MAX_CELLS
    assert n_rows <= tabular.MAX_ROWS
    assert n_cols * n_rows > tabular.MAX_CELLS
    header = ",".join(f"c{i}" for i in range(n_cols))
    row = ",".join("v" for _ in range(n_cols))
    too_many_cells = header + "\n" + "\n".join(row for _ in range(n_rows)) + "\n"
    with pytest.raises(tabular.TabularError) as cells_err:
        tabular.read_csv(too_many_cells)
    assert "cells" in str(cells_err.value).lower()


def test_number_and_date_stringification():
    import datetime

    data = _xlsx_bytes({"Nums": [
        ["int", "float", "whole", "date", "stamp"],
        [5, 2.5, 3.0, datetime.date(2026, 1, 2), datetime.datetime(2026, 3, 4, 9, 30, 0)],
    ]})
    cols, rows = tabular.read_xlsx(data, sheet=None)
    assert rows == [{
        "int": "5",
        "float": "2.5",
        "whole": "3",
        "date": "2026-01-02",
        "stamp": "2026-03-04T09:30:00",
    }]
