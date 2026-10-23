# Copyright (c) 2026 Noah Severs
# SPDX-License-Identifier: BUSL-1.1
"""Export-set sources: a set of report and list exports read as one company's books.

The fixture is a fictitious company's three years of trading, written the way the
source's export menus write it."""

from __future__ import annotations

import io
import zipfile
from datetime import date
from decimal import Decimal
from functools import cache

import openpyxl
import pytest

from celerp.importers.adapters.base import MigrationDecisions, ScanError
from celerp.importers.schema import CIFMode, CoverageClass, ReconciliationMeasure as M
from fixtures.manager_io.support import BASIC, artifact
from fixtures.qbo_export import build as qbo
from fixtures.xero_export import build
from migration_support import maker, real_engine  # noqa: F401
from test_migration_e2e import _maps, _passing, migrate_set
from ui.i18n import t

FULL = MigrationDecisions(mode=CIFMode.FULL_HISTORY)
CUTOVER = date(2025, 6, 30)


@cache
def _xero() -> tuple[tuple[str, bytes], ...]:
    return tuple(build.files().items())


@cache
def _qbo() -> tuple[tuple[str, bytes], ...]:
    return tuple(qbo.files().items())


def _set(tmp_path, files) -> list:
    folder = tmp_path / "upload"
    folder.mkdir(exist_ok=True)
    out = []
    for name, data in files:
        (folder / name).write_bytes(data)
        out.append(artifact(folder / name))
    return out


def _adapter():
    from celerp.importers.adapters.exports import XeroAdapter

    return XeroAdapter()


def _qbo_adapter():
    from celerp.importers.adapters.exports import QuickBooksOnlineAdapter

    return QuickBooksOnlineAdapter()


def _without(*names: str) -> list[tuple[str, bytes]]:
    return [(n, d) for n, d in _xero() if n not in names]


def _dated(data: bytes, heading: str) -> bytes:
    book = openpyxl.load_workbook(io.BytesIO(data))
    book.active["A3"] = heading
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# ── Recognition ───────────────────────────────────────────────────────────────

def test_xero_export_set_is_recognised_by_content_alone(tmp_path):
    from celerp.importers.adapters.registry import detect_adapter

    files = _set(tmp_path, [(f"export-{n}{name[name.rfind('.'):]}", data) for n, (name, data) in enumerate(_xero())])
    assert detect_adapter(files).key == "xero"
    scan = _adapter().inspect(files)
    assert scan.company_name == build.COMPANY
    assert scan.base_currency == "GBP"
    assert scan.currencies == ("GBP", "USD")
    assert (scan.period_start, scan.period_end) == (date(2023, 1, 1), date(2025, 12, 31))
    assert "foreign_currency" in scan.features

    # The Manager adapter does not claim a set, and the export adapter does not claim a Manager file.
    assert _adapter().detect([artifact(BASIC)]).matched is False
    assert detect_adapter([artifact(BASIC)]).key == "manager_io"


def test_xero_export_set_in_one_zip_reads_the_same(tmp_path):
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as z:
        for name, data in _xero():
            z.writestr(f"Harbour/{name}", data)
    zipped = _adapter().inspect(_set(tmp_path, [("xero-export.zip", archive.getvalue())]))
    loose = _adapter().inspect(_set(tmp_path, _xero()))
    assert zipped == loose


def test_a_set_missing_a_file_or_holding_a_stranger_names_each_file(tmp_path):
    files = _set(tmp_path, [*_without("General Ledger.txt"), ("notes.txt", b"call the bank\n")])
    assert _adapter().detect(files).matched is True
    with pytest.raises(ScanError) as caught:
        _adapter().inspect(files)
    message = str(caught.value)
    assert "General Ledger" in message                     # missing
    assert "notes.txt" in message                          # not recognised
    assert "ChartOfAccounts.csv (Chart of accounts)" in message   # what was recognised


def test_an_unreadable_file_is_named_and_not_also_called_missing(tmp_path):
    files = _set(tmp_path, [*_without("Contacts.csv"), ("Contacts.csv", b"\x81\x8d\x8f\x90\x9d" * 20)])
    with pytest.raises(ScanError) as caught:
        _adapter().inspect(files)
    message = str(caught.value)
    assert "Contacts.csv could not be read. Export it again and upload the new file." in message
    assert "? could not be read" not in message
    assert "Contacts is missing" not in message


def _many_files(count: int) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as z:
        for k in range(count):
            z.writestr(f"{k:08d}.csv", b"")
    return out.getvalue()


@pytest.mark.asyncio
@pytest.mark.parametrize("source", [None, "xero", "quickbooks_online"])
async def test_a_zip_listing_more_files_than_the_server_holds_is_refused_without_reading_it(client, monkeypatch,
                                                                                            tmp_path, source):
    """A zip of 300,000 empty files, offered to every reader or to one chosen: each refuses it from
    the size its directory declares, without holding the upload or its directory."""
    from celerp.config import settings
    from migration_support import recognition_peak, scan_upload

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    data = _many_files(300_000)
    peaks = recognition_peak(monkeypatch)
    r = await scan_upload(client, data, name="books.zip", source=source)
    assert r.status_code == 422, r.text
    assert peaks and max(peaks) < 16 * 1024 ** 2, f"{len(data) / 2 ** 20:.0f} MiB upload, peaks {peaks}"


def test_a_file_beyond_what_an_export_set_reads_is_named_without_being_read(tmp_path, monkeypatch):
    from celerp.importers.adapters.exports import files

    many = _set(tmp_path, [("many.zip", _many_files(300_000))])
    _, strangers = files.read_set(many, ())
    assert [s.reason for s in strangers] == [t("migration.err_export_archive_too_large", name="many.zip")]
    whole = _set(tmp_path, [("Contacts.csv", dict(_xero())["Contacts.csv"])])
    monkeypatch.setattr(files, "LIMITS", files.tabular.Limits(files.LIMITS.rows, files.LIMITS.cells,
                                                              whole[0].size_bytes - 1, files.LIMITS.uncompressed))
    monkeypatch.setattr(type(whole[0].path), "read_bytes", lambda self: pytest.fail(f"{self.name} was read"))
    _, strangers = files.read_set(whole, ())
    assert [s.reason for s in strangers] == [t("migration.err_file_too_large")]


def test_a_workbook_listing_more_entries_than_the_server_holds_is_refused_before_it_is_parsed():
    import tracemalloc

    from celerp.importers import tabular
    from celerp.importers.adapters.exports import files

    data = _many_files(300_000)
    tracemalloc.start()
    try:
        with pytest.raises(tabular.TabularError) as refused:
            tabular.workbook_lines(data, files.LIMITS)
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert str(refused.value) == t("import.err_xlsx_entries")
    assert peak < 16 * 1024 ** 2, f"{peak / 2 ** 20:.0f} MiB peak"


def test_overlapping_batches_are_refused(tmp_path):
    first = dict(_xero())["Invoices 1.csv"]
    with pytest.raises(ScanError, match="Invoices 1.csv.*Invoices 9.csv"):
        _adapter().inspect(_set(tmp_path, [*_xero(), ("Invoices 9.csv", first)]))


def test_a_year_without_its_trial_balance_is_refused(tmp_path):
    with pytest.raises(ScanError, match="Trial Balance.*2024"):
        _adapter().inspect(_set(tmp_path, _without("Trial Balance 2024.xlsx")))


def test_an_aged_report_older_than_the_ledger_is_refused(tmp_path):
    files = dict(_xero())
    files["Aged Receivables Detail.xlsx"] = _dated(files["Aged Receivables Detail.xlsx"], "As at 30 Nov 2025")
    with pytest.raises(ScanError, match="Aged Receivables Detail.*31 Dec 2025"):
        _adapter().inspect(_set(tmp_path, files.items()))


def test_a_ledger_account_missing_from_the_chart_is_refused(tmp_path):
    files = dict(_xero())
    chart = files["ChartOfAccounts.csv"].decode("utf-8-sig").splitlines()
    files["ChartOfAccounts.csv"] = "\r\n".join(ln for ln in chart if not ln.startswith("404,")).encode()
    with pytest.raises(ScanError, match="404 Bank Fees.*archived"):
        _adapter().inspect(_set(tmp_path, files.items()))


def test_slashed_dates_are_read_in_the_one_order_the_set_allows():
    from celerp.importers.adapters.exports.files import DateReader

    assert DateReader(["13/02/2024", "01/03/2024"])("01/03/2024") == date(2024, 3, 1)
    assert DateReader(["02/13/2024", "01/03/2024"])("01/03/2024") == date(2024, 1, 3)
    assert DateReader(["2024-03-01"])("2024-03-01") == date(2024, 3, 1)
    with pytest.raises(ScanError, match="day or the month"):
        DateReader(["01/02/2024", "03/04/2024"])
    with pytest.raises(ScanError, match="both"):
        DateReader(["13/02/2024", "02/13/2024"])


# ── Masters and documents ─────────────────────────────────────────────────────

def test_an_export_set_holds_no_attachment(tmp_path):
    with pytest.raises(ScanError, match="This attachment is not in the uploaded file."):
        _adapter().read_attachment(_set(tmp_path, _xero()), FULL, "attachment:1")


def test_xero_masters_and_documents_carry_into_cif(tmp_path):
    manifest = _adapter().build_manifest(_set(tmp_path, _xero()), FULL)
    bundle = manifest.bundle
    assert bundle.company.name == build.COMPANY and bundle.company.base_currency == "GBP"
    assert [c.code for c in bundle.currencies] == ["GBP", "USD"]

    accounts = {a.code: a for a in bundle.accounts}
    assert set(accounts) == {code for code, *_ in build.CHART}
    assert (accounts["091"].control, accounts["091"].currency) == ("bank", "USD")
    assert (accounts["610"].control, accounts["800"].control, accounts["820"].control) == (
        "receivable", "payable", "tax")
    assert accounts["499"].control == "exchange_difference"
    assert accounts["960"].control == "retained_earnings"
    assert (accounts["310"].account_type, accounts["416"].account_type, accounts["720"].account_type) == (
        "cogs", "expense", "asset")

    taxes = {tc.name: tc.rate_percent for tc in bundle.tax_codes}
    assert taxes["20% (VAT on Income)"] == Decimal(20) and taxes["Zero Rated Income"] == 0

    contacts = {c.name: c for c in bundle.contacts}
    assert set(contacts) >= {name for name, *_ in build.CUSTOMERS} | {name for name, _ in build.SUPPLIERS}
    assert contacts["Keystone Imports Inc"].currency == "USD"
    assert contacts["Keystone Imports Inc"].roles == ["customer"]
    assert contacts["Mill Timber Supply"].roles == ["supplier"]
    assert {i.sku for i in bundle.items} == {code for code, *_ in build.ITEMS}

    docs = {d.source_external_id: d for d in bundle.documents}
    source = {d.number: d for d in build.simulate().documents}
    assert len(docs) == len([d for d in source.values() if d.status != "Deleted"])
    statuses = {d.ref: d.status for d in docs.values()}
    assert [statuses[n] for n, d in source.items() if d.status == "Voided"] == ["void"] * 6
    assert [statuses[n] for n, d in source.items() if d.status == "Draft"] == ["draft"]
    usd = next(d for d in docs.values() if d.currency == "USD" and d.status == "paid")
    assert usd.exchange_rate is not None and usd.line_items[0].taxes[0].tax_code_external_id
    notes = [d for d in docs.values() if d.doc_type == "credit_note"]
    assert len(notes) == 3 and all(n.metadata["applications"] == {f"S:{source[n.ref].reference}": str(n.total)} for n in notes)
    assert docs["S:INV-0001"].line_items[0].item_external_id is not None
    assert [ln.tracking for ln in docs["S:INV-0001"].line_items] == [{"Region": "North", "Project": "Beta"}]

    coverage = {row.source_type: (row.count, row.coverage_class) for row in manifest.coverage}
    assert coverage["Deleted invoice or bill"] == (3, CoverageClass.IGNORED_NON_BUSINESS)
    assert coverage["Tracking category"][1] == CoverageClass.UNSUPPORTED_NONFINANCIAL
    # The report's note is read from the catalogs, so it is shown in the user's language.
    notes = {row.source_type: row.note for row in manifest.coverage}
    assert notes["Tracking category"] == t("migration.coverage_tracking_note") != "migration.coverage_tracking_note"
    assert all(row.coverage_class != CoverageClass.UNSUPPORTED_FINANCIAL_BLOCKER for row in manifest.coverage)


def test_xero_expectations_come_from_the_reports_not_the_conversion(tmp_path):
    expectations = _adapter().source_expectations(_set(tmp_path, _xero()), FULL).expectations
    by = {(e.measure, e.key): e.expected for e in expectations}
    books = build.simulate()
    end = books.balances(build.END)
    accounts = {a: a for a, *_ in build.CHART}
    # The trial balance at the last month end, and at every month end before it.
    assert by[(M.TRIAL_BALANCE, accounts["090"])] == end["090"]
    assert by[(M.BANK_CASH, accounts["091"])] == end["091"]
    june = books.balances(date(2024, 6, 30))
    assert by[(M.TRIAL_BALANCE_AT, "090@2024-06-30")] == june["090"]
    assert len({k for m, k in by if m == M.TRIAL_BALANCE_AT and k.endswith("@2023-01-31")}) > 5
    # Customer balances from the aged receivables report.
    assert by[(M.AR_BY_CUSTOMER, "Lakeshore Trading Co")] == sum(
        (p.amount for j in books.journals if j.contact == "Lakeshore Trading Co" for p in j.postings
         if p.account == "610"), Decimal(0))


# ── QuickBooks Online ─────────────────────────────────────────────────────────

def test_quickbooks_online_export_set_is_recognised_by_content_alone(tmp_path):
    from celerp.importers.adapters.registry import detect_adapter

    files = _set(tmp_path, [(f"report-{n}.xlsx", data) for n, (_, data) in enumerate(_qbo())])
    assert detect_adapter(files).key == "quickbooks_online"
    scan = _qbo_adapter().inspect(files)
    assert scan.company_name == build.COMPANY
    assert scan.base_currency == "GBP"
    assert scan.currencies == ("GBP", "USD")
    assert (scan.period_start, scan.period_end) == (date(2023, 1, 1), date(2025, 12, 31))
    assert "foreign_currency" in scan.features
    # Neither export-set adapter claims the other's set.
    assert _adapter().detect(files).matched is False
    assert _qbo_adapter().detect(_set(tmp_path, _xero())).matched is False


def test_quickbooks_online_export_set_in_one_zip_reads_the_same(tmp_path):
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as z:
        for name, data in _qbo():
            z.writestr(f"Harbour/{name}", data)
    zipped = _qbo_adapter().inspect(_set(tmp_path, [("quickbooks-export.zip", archive.getvalue())]))
    assert zipped == _qbo_adapter().inspect(_set(tmp_path, _qbo()))


def test_a_quickbooks_online_set_missing_its_journal_names_it(tmp_path):
    with pytest.raises(ScanError, match="Journal") as caught:
        _qbo_adapter().inspect(_set(tmp_path, [(n, d) for n, d in _qbo() if n != "Journal.xlsx"]))
    assert "Account List.xlsx (Account List)" in str(caught.value)


def test_a_quickbooks_online_year_without_its_closing_trial_balance_is_refused(tmp_path):
    files = [(n, d) for n, d in _qbo() if n != "Trial Balance 2024-12-31.xlsx"]
    with pytest.raises(ScanError, match="Trial Balance.*31 Dec 2024"):
        _qbo_adapter().inspect(_set(tmp_path, files))


def test_quickbooks_online_masters_and_documents_carry_into_cif(tmp_path):
    manifest = _qbo_adapter().build_manifest(_set(tmp_path, _qbo()), FULL)
    bundle = manifest.bundle
    assert bundle.company.name == build.COMPANY and bundle.company.base_currency == "GBP"
    accounts = {a.code: a for a in bundle.accounts}
    assert set(accounts) == {code for code, *_ in build.CHART}
    assert accounts["499"].name == "Exchange Gain or Loss"
    assert (accounts["091"].control, accounts["091"].currency) == ("bank", "USD")
    assert (accounts["610"].control, accounts["800"].control, accounts["820"].control) == (
        "receivable", "payable", "tax")
    assert (accounts["499"].control, accounts["960"].control) == ("exchange_difference", "retained_earnings")
    assert (accounts["270"].account_type, accounts["310"].account_type, accounts["721"].account_type) == (
        "revenue", "cogs", "asset")

    taxes = {tc.name: tc.rate_percent for tc in bundle.tax_codes}
    assert taxes["20.0% S"] == Decimal(20) and taxes["0.0% Z"] == 0

    contacts = {c.name: c for c in bundle.contacts}
    assert contacts["Keystone Imports Inc"].currency == "USD"
    assert contacts["Keystone Imports Inc"].roles == ["customer"]
    assert contacts["Mill Timber Supply"].roles == ["supplier"]
    assert {i.sku for i in bundle.items} == {code for code, *_ in build.ITEMS}

    docs = {d.source_external_id: d for d in bundle.documents}
    source = {d.number: d for d in build.simulate().documents if d.status not in ("Deleted", "Draft")}
    assert set(docs) == {f"{'S' if d.sales else 'P'}:{n}" for n, d in source.items()}
    assert [docs[f"S:{n}"].status for n, d in source.items() if d.status == "Voided"] == ["void"] * 6
    usd = next(d for d in docs.values() if d.currency == "USD" and d.status == "paid")
    assert usd.exchange_rate is not None and usd.line_items[0].taxes[0].tax_code_external_id
    notes = [d for d in docs.values() if d.doc_type == "credit_note"]
    assert len(notes) == 3 and all(n.metadata["applications"] == {f"S:{source[n.ref].reference}": str(n.total)} for n in notes)
    assert docs["S:INV-0001"].line_items[0].item_external_id is not None
    assert [ln.tracking for ln in docs["S:INV-0001"].line_items] == [{"Class": "North", "Location": "Beta"}]

    # Payments are matched to what they settled, so few stay plain journals.
    settled = [s for s in bundle.settlements if s.settlement_type == "receipt"]
    assert len(settled) > 100
    assert all(row.coverage_class != CoverageClass.UNSUPPORTED_FINANCIAL_BLOCKER for row in manifest.coverage)


def test_a_quickbooks_online_payment_applied_to_two_invoices_settles_each(tmp_path):
    """One receipt applied to two invoices of a customer settles each by what was applied to it."""
    books = build.simulate()
    pays = [j for j in books.journals if j.source == "Receivable Payment" and j.currency == build.BASE]
    first, second = next((a, b) for a in pays for b in pays if a.contact == b.contact and a.invoice < b.invoice)
    party = sum((p.amount for p in first.postings if p.account == "610"), Decimal(0))
    later = sum((p.amount for p in second.postings if p.account == "610"), Decimal(0))
    first.applied = {first.invoice: -party, second.invoice: -later}
    first.on, first.invoice = max(first.on, second.on), ""
    totals: dict[str, Decimal] = {}
    for p in first.postings + second.postings:
        totals[p.account] = totals.get(p.account, Decimal(0)) + p.amount
    first.postings = [build.Posting(account, amount) for account, amount in totals.items()]
    books.journals.remove(second)
    bundle = _qbo_adapter().build_manifest(_set(tmp_path, qbo.files(books).items()), FULL).bundle
    settled = {tuple((a.document_external_id, a.amount) for a in s.allocations) for s in bundle.settlements}
    assert ((f"S:{next(iter(first.applied))}", -party), (f"S:{list(first.applied)[1]}", -later)) in settled


def test_quickbooks_online_expectations_come_from_the_reports(tmp_path):
    expectations = _qbo_adapter().source_expectations(_set(tmp_path, _qbo()), FULL).expectations
    by = {(e.measure, e.key): e.expected for e in expectations}
    books = build.simulate()
    end = books.balances(build.END)
    assert by[(M.TRIAL_BALANCE, "090")] == end["090"]
    assert by[(M.BANK_CASH, "091")] == end["091"]
    assert by[(M.TRIAL_BALANCE_AT, "090@2024-06-30")] == books.balances(date(2024, 6, 30))["090"]
    assert by[(M.TRIAL_BALANCE_AT, "200@2024-06-30")] == books.balances(date(2024, 6, 30))["200"]
    assert by[(M.AR_BY_CUSTOMER, "Lakeshore Trading Co")] == sum(
        (p.amount for j in books.journals if j.contact == "Lakeshore Trading Co" for p in j.postings
         if p.account == "610"), Decimal(0))


# ── End to end ────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
@pytest.mark.timeout(180)  # three years of trading through the real runner
@pytest.mark.parametrize("decisions", [{"mode": "full_history"},
                                       {"mode": "cutover", "cutover_date": CUTOVER.isoformat()}])
@pytest.mark.parametrize("source", ["quickbooks_online", "xero"])
async def test_quickbooks_and_xero_adapters(real_engine, monkeypatch, tmp_path, decisions, source):
    """An export set migrates through the real runner with nothing rejected and every
    reconciliation check passing, month-end trial balances included; every fallback is
    listed as mapped with loss."""
    from celerp.services import migrations

    files = list(_qbo() if source == "quickbooks_online" else _xero())
    run, rejected = await migrate_set(real_engine, files, decisions, monkeypatch, tmp_path)
    assert rejected == []
    _passing(run)
    measures = {r["check"] for r in run.reconciliation["rows"]}
    assert {"trial_balance_at", "ar_by_customer", "ap_by_supplier", "bank_cash"} <= measures
    async with maker(real_engine)() as s:
        report = await migrations.run_view(s, run)
    coverage = {(e["source_type"], e["coverage_class"]) for e in report["coverage"]}
    for m in await _maps(real_engine, run):
        if m.meta.get("representation") == "journal_fallback":
            assert ((m.source_type, "mapped_with_loss") in coverage
                    or (f"{m.source_type} (journal fallback)", "mapped_with_loss") in coverage), m.source_type


def test_a_document_number_is_found_only_on_its_own_contacts_document():
    from celerp.importers.adapters.exports.book import Book, Document
    from celerp.importers.adapters.exports.ledger import find_document

    def doc(key: str, contact: str) -> Document:
        return Document(key, "INV-1", "invoice", "Invoice", contact, "", date(2025, 1, 1), None, "USD", "posted",
                        [], Decimal(0), Decimal(0), Decimal(0), Decimal(0))

    book = Book("xero")
    book.documents = {k: doc(k, c) for k, c in (("S:INV-1", "Ann"), ("S:INV-1|Bob", "Bob"), ("P:INV-1|Cy", "Cy"))}
    assert find_document(book, True, "INV-1", "Ann").key == "S:INV-1"
    assert find_document(book, True, "INV-1", "Bob").key == "S:INV-1|Bob"
    assert find_document(book, True, "INV-1", "Dee") is None          # never another contact's
    assert find_document(book, True, "INV-1", "").key == "S:INV-1"    # a journal naming no contact
    assert find_document(book, False, "INV-1", "Cy").key == "P:INV-1|Cy"
    assert find_document(book, False, "INV-1", "Ann") is None
    # Two of Bob's documents share the number: a reference to it names neither, and is refused.
    book.repeated.add(("S", "INV-1", "Bob"))
    with pytest.raises(ScanError) as refused:
        find_document(book, True, "INV-1", "Bob")
    assert str(refused.value) == t("migration.err_export_reference_ambiguous", number="INV-1", contact="Bob")
    assert find_document(book, True, "INV-1", "Ann").key == "S:INV-1"


# ── Values an export set holds ────────────────────────────────────────────────

@pytest.mark.parametrize("text, expected", [
    ("1,234.50", Decimal("1234.50")), ("$1,234.50", Decimal("1234.50")), ("-$5.00", Decimal("-5.00")),
    ("Br 1,234.50", Decimal("1234.50")), ("1,234.50 Br", Decimal("1234.50")), ("ETB 10", Decimal("10")),
    ("(Br 2.00)", Decimal("-2.00")), ("ብር 3.25", Decimal("3.25")), ("1e-05", Decimal("0.00001")), ("", Decimal(0)),
])
def test_an_amount_is_read_with_its_currency_sign_and_separators(text, expected):
    from celerp.importers.adapters.exports.files import amount

    assert amount(text) == expected


@pytest.mark.parametrize("text", ["NaN", "Infinity", "-inf", "1e999999", "9" * 40, "=1+1", "12..5", "Br", "1 2"])
def test_an_amount_that_is_not_a_plain_number_in_range_is_refused_by_name(text):
    from celerp.importers.adapters.exports.files import amount, money

    for read in (amount, money):
        with pytest.raises(ScanError, match="is not a number"):
            read(text, "Sales.csv INV-1")


def test_dates_with_a_two_digit_year_read_in_the_order_the_source_writes():
    from celerp.importers.adapters.exports.files import DateReader

    month_first = DateReader(["03/04/25", "01/02/2025"], day_first=False)
    assert month_first("03/04/25") == date(2025, 3, 4)
    assert month_first("12/31/99") == date(1999, 12, 31)
    with pytest.raises(ScanError, match="day or the month"):
        DateReader(["03/04/25"])
    with pytest.raises(ScanError, match="is not a date"):
        month_first("02/30/25", "General.csv")


def test_a_cell_longer_than_the_reader_takes_names_its_file_as_unreadable(tmp_path):
    from celerp.importers.adapters.exports import files

    huge = b"Code,Name\r\n200,\"" + b"x" * (1024 * 1024) + b"\"\r\n"
    _, strangers = files.read_set(_set(tmp_path, [("Chart.csv", huge)]), ())
    assert [s.name for s in strangers] == ["Chart.csv"]
    assert strangers[0].reason.startswith("Chart.csv could not be read")


def test_a_delimited_file_is_refused_at_its_row_bound_without_holding_every_row():
    import tracemalloc

    from celerp.importers import tabular

    text = "a,b,c\r\n" * 400_000
    tracemalloc.start()
    try:
        with pytest.raises(tabular.TabularError):
            tabular.delimited_lines(text, ",", tabular.Limits(rows=1000, cells=10_000, file_bytes=10 ** 9,
                                                              uncompressed=10 ** 9))
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert peak < 4 * 1024 ** 2, f"{peak / 2 ** 20:.1f} MiB peak"


@pytest.mark.parametrize("error", [ValueError("bad"), KeyError("x"), TypeError("y"), ArithmeticError("z")])
def test_a_value_an_export_set_reader_cannot_convert_refuses_the_set_as_damaged(tmp_path, error):
    from celerp.importers.adapters.base import SourceRevisionError

    adapter = _adapter()

    def broken(artifacts):
        raise error

    adapter.read = broken
    files = _set(tmp_path, _xero())
    for step in (lambda: adapter.inspect(files), lambda: adapter.build_manifest(files, FULL),
                 lambda: adapter.source_expectations(files, FULL)):
        with pytest.raises(SourceRevisionError) as refused:
            step()
        assert str(refused.value) == t("migration.err_export_damaged", source="Xero")


def _broken_workbook() -> bytes:
    """A workbook whose sheet is not well-formed XML: a file damaged in copying."""
    book = openpyxl.Workbook()
    book.active.append(["Account", "Debit"])
    saved = io.BytesIO()
    book.save(saved)
    source = zipfile.ZipFile(io.BytesIO(saved.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for info in source.infolist():
            data = source.read(info)
            z.writestr(info, b"<worksheet><sheetData><row><<<" if info.filename == "xl/worksheets/sheet1.xml" else data)
    return out.getvalue()


def test_a_damaged_workbook_is_named_as_unreadable(tmp_path):
    from celerp.importers.adapters.exports import files

    _, strangers = files.read_set(_set(tmp_path, [("Trial Balance.xlsx", _broken_workbook())]), ())
    assert [s.reason for s in strangers] == [t("migration.err_export_unreadable", name="Trial Balance.xlsx")]


# ── Headers and amounts, as every export set reads them ──────────────────────

def _xero_edited(name: str, old: str, new: str) -> list[tuple[str, bytes]]:
    """The Xero set with `old` written as `new` once in one of its text files."""
    out = []
    for n, data in _xero():
        if n == name:
            text = data.decode("utf-8-sig")
            assert old in text, old
            data = text.replace(old, new, 1).encode("utf-8-sig")
        out.append((n, data))
    return out


def _qbo_cell(name: str, coordinate: str, value: str) -> list[tuple[str, bytes]]:
    """The QuickBooks Online set with one cell of one workbook set to `value`."""
    out = []
    for n, data in _qbo():
        if n == name:
            book = openpyxl.load_workbook(io.BytesIO(data))
            book.active[coordinate] = value
            saved = io.BytesIO()
            book.save(saved)
            data = saved.getvalue()
        out.append((n, data))
    return out


def _refused(adapter, tmp_path, files) -> str:
    with pytest.raises(ScanError) as refused:
        adapter.inspect(_set(tmp_path, files))
    return str(refused.value)


def test_a_column_headed_twice_is_refused_rather_than_one_value_replacing_the_other(tmp_path):
    # "Description" headed "Name" again: *Name and Name are one column to the reader.
    files = _xero_edited("ChartOfAccounts.csv", "*Tax Code,Description,", "*Tax Code,Name,")
    assert t("migration.err_export_column_twice", name="ChartOfAccounts.csv", column="Name") in _refused(
        _adapter(), tmp_path, files)
    files = _qbo_cell("Account List.xlsx", "E5", "Account")
    assert t("migration.err_export_column_twice", name="Account List.xlsx", column="Account") in _refused(
        _qbo_adapter(), tmp_path, files)


def test_a_value_under_no_heading_is_refused_rather_than_dropped(tmp_path):
    files = _xero_edited("ChartOfAccounts.csv", "090,Business Bank Account,Bank,No VAT,,No,No,No,",
                         "090,Business Bank Account,Bank,No VAT,,No,No,No,,stray")
    assert t("migration.err_export_value_unheaded", name="ChartOfAccounts.csv", value="stray") in _refused(
        _adapter(), tmp_path, files)
    files = _qbo_cell("Account List.xlsx", "G6", "stray")
    assert t("migration.err_export_value_unheaded", name="Account List.xlsx", value="stray") in _refused(
        _qbo_adapter(), tmp_path, files)


def test_an_amount_is_read_exactly_as_printed_and_a_workbooks_binary_fraction_as_meant():
    from celerp.importers.adapters.exports.files import money

    assert money("1.234") == Decimal("1.234") and str(money("1.234")) == "1.234"
    assert money("1,234.5") == Decimal("1234.5")
    assert money("0.30000000000000004") == Decimal("0.3")
    assert money("1e3") == Decimal(1000) and str(money("1e3")) == "1000"


def test_an_amount_finer_than_its_currency_is_refused_never_rounded(tmp_path):
    """A pound amount of three decimal places cannot be the books' figure: it is refused by name,
    where reading it at cents would have changed it unseen."""
    files = _xero_edited("Invoices 1.csv", "INV-0001,,01/01/2023,31/01/2023,10.08,1.68,10.08,0.00",
                         "INV-0001,,01/01/2023,31/01/2023,10.08,1.68,10.075,0.005")
    assert _refused(_adapter(), tmp_path, files) == t(
        "migration.err_export_precision", where="Invoices 1.csv INV-0001", value="10.075", currency="GBP", places=2)


def test_xero_and_quickbooks_sets_read_the_same_with_amounts_kept_exact(tmp_path):
    """Every amount both fixtures print fits its currency, so reading them exactly changes nothing."""
    for adapter, files in ((_adapter(), _xero()), (_qbo_adapter(), _qbo())):
        manifest = adapter.build_manifest(_set(tmp_path, files), FULL)
        for doc in manifest.bundle.documents:
            assert doc.total == doc.total.quantize(Decimal("0.01")), doc.source_external_id
        for journal in manifest.bundle.journals:
            for line in journal.lines:
                assert line.debit == line.debit.quantize(Decimal("0.01")), journal.source_external_id
